import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from backend.research.schemas import Pass3Output


def build_docx(
    query: str,
    depth: int,
    pass_3: Pass3Output,
    include_sources: bool = True,
    include_open_questions: bool = True,
) -> io.BytesIO:
    """Generate a DOCX report from Pass 3 output."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # ── Cover page ───────────────────────────────────────────────
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Deep Research Report")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Query: {query}\n").bold = True
    meta.add_run(f"Date: {datetime.utcnow().strftime('%B %d, %Y')}\n")
    meta.add_run(f"Research Depth: Level {depth}\n")
    meta.add_run(
        f"Overall Confidence: {pass_3.metadata.research_confidence_overall.upper()}"
    )

    doc.add_page_break()

    # ── Executive Summary ────────────────────────────────────────
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(pass_3.report.executive_summary)
    doc.add_paragraph()

    # ── Report Sections ──────────────────────────────────────────
    for section in pass_3.report.sections:
        doc.add_heading(section.title, level=2)
        for line in section.content.split("\n"):
            stripped = line.strip()
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                para = doc.add_paragraph(stripped[2:], style="List Bullet")
                _apply_confidence_if_needed(para)
            elif stripped:
                doc.add_paragraph(stripped)

        conf_para = doc.add_paragraph()
        conf_run = conf_para.add_run(
            f"Confidence: {section.confidence.upper()} — {section.confidence_rationale}"
        )
        conf_run.italic = True
        conf_run.font.size = Pt(9)
        color = _confidence_color(section.confidence)
        conf_run.font.color.rgb = color

    # ── Key Conclusions ──────────────────────────────────────────
    doc.add_heading("Key Conclusions", level=1)
    for kc in pass_3.report.key_conclusions:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(kc.conclusion)
        run.bold = True
        conf_run = para.add_run(f"  [{kc.confidence.upper()}]")
        conf_run.font.color.rgb = _confidence_color(kc.confidence)
        conf_run.font.size = Pt(9)

    # ── Open Questions ───────────────────────────────────────────
    if include_open_questions and pass_3.report.open_questions:
        doc.add_heading("Open Questions", level=1)
        for oq in pass_3.report.open_questions:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(oq.question).bold = True
            para.add_run(f"\n  Why it matters: {oq.why_it_matters}")

    # ── Sources Appendix ─────────────────────────────────────────
    if include_sources and pass_3.sources:
        doc.add_heading("Sources", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        headers = table.rows[0].cells
        headers[0].text = "#"
        headers[1].text = "Title"
        headers[2].text = "URL"
        headers[3].text = "Pass"

        for i, src in enumerate(pass_3.sources, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = src.title
            row[2].text = src.url
            row[3].text = str(src.found_in_pass)

    # ── Save to buffer ───────────────────────────────────────────
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_markdown(
    query: str,
    depth: int,
    pass_3: Pass3Output,
    include_sources: bool = True,
    include_open_questions: bool = True,
) -> str:
    """Generate a Markdown report from Pass 3 output."""
    lines: list[str] = []
    lines.append(f"# Deep Research Report: {query}")
    lines.append("")
    lines.append(f"**Date:** {datetime.utcnow().strftime('%B %d, %Y')}")
    lines.append(f"**Depth:** Level {depth}")
    lines.append(
        f"**Overall Confidence:** {pass_3.metadata.research_confidence_overall.upper()}"
    )
    lines.append("")
    lines.append("## Executive Summary")
    lines.append("")
    lines.append(pass_3.report.executive_summary)
    lines.append("")

    for section in pass_3.report.sections:
        lines.append(f"## {section.title}")
        lines.append("")
        lines.append(section.content)
        lines.append("")
        lines.append(
            f"*Confidence: {section.confidence.upper()} — {section.confidence_rationale}*"
        )
        lines.append("")

    lines.append("## Key Conclusions")
    lines.append("")
    for kc in pass_3.report.key_conclusions:
        lines.append(f"- **{kc.conclusion}** [{kc.confidence.upper()}]")
    lines.append("")

    if include_open_questions and pass_3.report.open_questions:
        lines.append("## Open Questions")
        lines.append("")
        for oq in pass_3.report.open_questions:
            lines.append(f"- **{oq.question}** — {oq.why_it_matters}")
        lines.append("")

    if include_sources and pass_3.sources:
        lines.append("## Sources")
        lines.append("")
        lines.append("| # | Title | URL | Pass |")
        lines.append("|---|-------|-----|------|")
        for i, src in enumerate(pass_3.sources, 1):
            lines.append(f"| {i} | {src.title.replace('|', '\\|')} | {src.url.replace('|', '\\|')} | {src.found_in_pass} |")
        lines.append("")

    return "\n".join(lines)


def _confidence_color(level: str) -> RGBColor:
    """Return color for confidence level."""
    if level == "high":
        return RGBColor(0x16, 0xA3, 0x4A)
    elif level == "medium":
        return RGBColor(0xD9, 0x77, 0x06)
    return RGBColor(0xDC, 0x26, 0x26)


def _apply_confidence_if_needed(para) -> None:  # type: ignore[no-untyped-def]
    """Placeholder for future confidence styling on paragraphs."""
    pass
